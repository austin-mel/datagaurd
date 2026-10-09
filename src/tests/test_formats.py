import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from fastapi.testclient import TestClient

from src.errors import InputError
from src.main import create_app
from src.services.tabular import MEDIA_TYPES, parse_dataset, replace_values
from src.settings import Settings
from src.tests.test_remediation import approve, propose


def encode(file_format: str, headers: list[str], rows: list[list[Any]], *, extra_sheet: bool = False) -> bytes:
    buffer = io.BytesIO()
    if file_format == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.title = "Records"
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        if extra_sheet:
            book.create_sheet("Notes")["A1"] = "Keep this sheet"
        book.save(buffer)
        book.close()
    elif file_format == "parquet":
        table = pa.table({header: [row[index] for row in rows] for index, header in enumerate(headers)})
        pq.write_table(table, buffer)
    else:
        text = io.StringIO(newline="")
        writer = csv.writer(text, delimiter="\t" if file_format == "tsv" else ",")
        writer.writerow(headers)
        writer.writerows(rows)
        buffer.write(text.getvalue().encode())
    return buffer.getvalue()


@pytest.mark.parametrize("file_format", list(MEDIA_TYPES))
def test_full_format_workflow(settings: Settings, file_format: str) -> None:
    content = encode(file_format, ["id", "amount"], [["001", "201"], ["002", "20"]], extra_sheet=file_format == "xlsx")
    form = {"sheet_name": "Records"} if file_format == "xlsx" else {}
    with TestClient(create_app(settings)) as client:
        rules = {"schema_version": 1, "name": "Orders", "structural_rule_id": "COLUMNS",
                 "required_columns": ["id", "amount"], "identifier_column": "id", "rules": [
                     {"id": "ID_UNIQUE", "kind": "unique", "column": "id", "explanation": "IDs are unique."},
                     {"id": "AMOUNT_NUMBER", "kind": "numeric", "column": "amount", "explanation": "Numeric amount."},
                     {"id": "AMOUNT_RANGE", "kind": "range", "column": "amount", "minimum": 0, "maximum": 100,
                      "prerequisite": "AMOUNT_NUMBER", "explanation": "Amount from 0 to 100."}]}
        rule_response = client.post("/rule-sets", json=rules)
        assert rule_response.status_code == 201, rule_response.text
        form["rule_set_id"] = rule_response.json()["id"]
        files = {"file": (f"orders.{file_format}", content, MEDIA_TYPES[file_format])}
        for endpoint in ("profile", "check"):
            result = client.post(f"/datasets/{endpoint}", files=files, data=form)
            assert result.status_code == 200, result.text
            assert result.json()["metadata"]["format"] == file_format
        response = client.post("/datasets/upload", files=files, data=form)
        assert response.status_code == 200, response.text
        saved = response.json()
        assert saved["check_run"]["processing_status"] == "completed"
        assert saved["check_run"]["validation_status"] == "failed"
        assert saved["check_run"]["rule_set_id"] == form["rule_set_id"]
        assert saved["check_run"]["result"]["metrics"]["duplicate_id_rate"]["percentage"] == 0
        finding = client.get(f"/check-runs/{saved['check_run']['id']}/findings").json()["items"][0]
        record = client.get(f"/findings/{finding['id']}/records").json()["items"][0]
        assert record["values"] == {"id": "001", "amount": "201"}
        identity = propose(client, saved, column="amount", records=[{"record_number": 1, "expected_value": "201"}], replacement="100")
        approve(client, identity)
        execution = client.post(f"/remediations/{identity}/execute")
        assert execution.status_code == 200, execution.text
        result = execution.json()
        assert result["check_run"]["validation_status"] == "passed"
        assert result["version"]["file_id"].endswith("." + file_format)
        downloaded = client.get(f"/versions/{result['version']['id']}/file")
        assert downloaded.headers["content-type"].split(";")[0] == MEDIA_TYPES[file_format]
        corrected = parse_dataset(downloaded.content, f"orders.{file_format}", settings, sheet_name=form.get("sheet_name"))
        assert corrected.records == (("001", "100"), ("002", "20"))
        assert client.get(f"/versions/{saved['version']['id']}/file").content == content
        if file_format == "xlsx":
            book = openpyxl.load_workbook(io.BytesIO(downloaded.content))
            assert book["Notes"]["A1"].value == "Keep this sheet"
            book.close()
    with TestClient(create_app(settings)) as client:
        assert client.post(f"/remediations/{identity}/execute").json() == result
        assert client.post(f"/versions/{result['version']['id']}/check").json()["validation_status"] == "passed"


def test_tsv_quoting_and_bom(settings: Settings) -> None:
    content = b'\xef\xbb\xbfid\tnote\n001\t"tab\tand\nnewline"\n'
    assert parse_dataset(content, "x.TSV", settings).records == (("001", "tab\tand\nnewline"),)
    with pytest.raises(InputError, match="quoted"):
        parse_dataset(b'id\tnote\n1\t"x"bad', "x.tsv", settings)


def test_excel_sheet_selection_and_cells(client: TestClient, settings: Settings) -> None:
    content = encode("xlsx", ["id", "day", "active"], [["001", date(2026, 10, 7), True]], extra_sheet=True)
    files = {"file": ("x.xlsx", content, MEDIA_TYPES["xlsx"])}
    response = client.post("/datasets/upload", files=files)
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "sheet_required"
    assert client.post("/datasets/upload", files=files, data={"sheet_name": "Unknown"}).status_code == 422
    parsed = parse_dataset(content, "x.xlsx", settings, sheet_name="Records")
    assert parsed.records == (("001", "2026-10-07", "true"),)
    assert parsed.metadata.encoding is None
    for value in ("=SUM(1,2)", "#DIV/0!"):
        bad = encode("xlsx", ["value"], [[value]])
        with pytest.raises(InputError) as error:
            parse_dataset(bad, "x.xlsx", settings)
        assert error.value.code == "unsupported_cell"
    with pytest.raises(InputError):
        parse_dataset(b"id\n1", "x.csv", settings, sheet_name="Records")


def test_parquet_types_and_exact_corrections(settings: Settings) -> None:
    content = encode("parquet", ["id", "amount", "day", "active", "stamp", "price"], [
        ["001", 201, date(2026, 10, 7), True, datetime(2026, 10, 7, 12), Decimal("1.25")],
        ["002", None, None, False, None, None]])
    parsed = parse_dataset(content, "x.parquet", settings)
    assert parsed.records[0] == ("001", "201", "2026-10-07", "true", "2026-10-07T12:00:00", "1.25")
    assert parsed.records[1] == ("002", "", "", "false", "", "")
    corrected, result = replace_values(content, parsed, 1, {1}, "100", settings)
    assert result.records[0][1] == "100"
    original_table = pq.read_table(io.BytesIO(content))
    updated_table = pq.read_table(io.BytesIO(corrected))
    assert original_table.schema == updated_table.schema
    for column in (0, 2, 3, 4, 5):
        assert original_table.column(column).equals(updated_table.column(column))
    for replacement in ("001", "bad", "1.5"):
        with pytest.raises(InputError) as error:
            replace_values(content, parsed, 1, {1}, replacement, settings)
        assert error.value.code == "incompatible_replacement"
    with pytest.raises(InputError) as error:
        parse_dataset(encode("parquet", ["nested"], [[[1, 2]]]), "x.parquet", settings)
    assert error.value.code == "unsupported_column_type"


@pytest.mark.parametrize("file_format", list(MEDIA_TYPES))
@pytest.mark.parametrize(("limits", "code"), [({"max_rows": 1}, "too_many_rows"),
    ({"max_columns": 1}, "too_many_columns"), ({"max_cells": 3}, "too_many_cells"),
    ({"max_decoded_bytes": 1}, "decoded_too_large"), ({"max_upload_bytes": 1}, "upload_too_large")])
def test_format_limits(settings: Settings, file_format: str, limits: dict[str, int], code: str) -> None:
    content = encode(file_format, ["id", "amount"], [["001", "10"], ["002", "20"]])
    with pytest.raises(InputError) as error:
        parse_dataset(content, "x." + file_format, settings.model_copy(update=limits))
    assert error.value.code == code
    assert error.value.status_code == 413


@pytest.mark.parametrize("file_format", ["xlsx", "parquet"])
def test_malformed_binary_upload(client: TestClient, file_format: str) -> None:
    response = client.post("/datasets/upload", files={"file": ("x." + file_format, b"bad", MEDIA_TYPES[file_format])})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "malformed_file"
    assert client.get("/datasets").json()["total"] == 0


def test_mixed_format_versions(client: TestClient) -> None:
    saved = client.post("/datasets/upload", files={"file": ("x.csv", b"id\n001\n")}).json()
    previous = saved["version"]["id"]
    for file_format in ("tsv", "xlsx", "parquet"):
        content = encode(file_format, ["id"], [["001"]])
        response = client.post(f"/datasets/{saved['dataset']['id']}/versions",
                               files={"file": ("x." + file_format, content, MEDIA_TYPES[file_format])})
        assert response.status_code == 200, response.text
        result = response.json()
        assert result["check_run"]["processing_status"] == "completed"
        assert result["check_run"]["baseline_id"] == previous
        schema = next(item for item in result["check_run"]["result"]["analysis_results"] if item["check_id"] == "SCHEMA_COLUMNS")
        assert schema["status"] == "passed"
        previous = result["version"]["id"]
