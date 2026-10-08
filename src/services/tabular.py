import csv
import hashlib
import io
from collections.abc import Iterable
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import PurePath
from typing import Any, cast
from zipfile import ZipFile

import openpyxl
from openpyxl.styles.numbers import is_datetime
import pyarrow as pa
import pyarrow.parquet as pq

from src.errors import InputError
from src.schemas import DataFormat, DatasetMetadata
from src.settings import Settings

csv.field_size_limit(2**31 - 1)
MEDIA_TYPES = {
    "csv": "text/csv", "tsv": "text/tab-separated-values",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "parquet": "application/vnd.apache.parquet",
}
MEDIA_ALIASES = {"csv": {"application/csv", "application/vnd.ms-excel", "text/plain"},
                 "tsv": {"text/tsv", "text/plain"}, "xlsx": set(),
                 "parquet": {"application/x-parquet"}}


@dataclass(frozen=True)
class ParsedDataset:
    metadata: DatasetMetadata
    records: tuple[tuple[str, ...], ...]

    def column(self, name: str) -> list[str]:
        index = self.metadata.headers.index(name)
        return [row[index] for row in self.records]


def is_missing(value: str) -> bool:
    return not value.strip()


def validate_file_type(filename: str | None, content_type: str | None) -> str:
    extension = PurePath(filename or "").suffix.lower().lstrip(".")
    if not filename or extension not in MEDIA_TYPES:
        raise InputError("unsupported_file_type", "Upload a .csv, .tsv, .xlsx, or .parquet file.", 415)
    media_type = (content_type or "").split(";", 1)[0].strip().lower()
    if media_type and media_type not in {MEDIA_TYPES[extension], "application/octet-stream", *MEDIA_ALIASES[extension]}:
        raise InputError("unsupported_file_type", "The supplied media type does not match the file extension.", 415)
    return filename


def _check_quoting(text: str, delimiter: str) -> None:
    # csv.reader(strict=True) still accepts quotes inside unquoted fields.
    state = "start"
    for char in text:
        if state == "quoted":
            if char == '"':
                state = "after_quote"
        elif state == "after_quote":
            if char == '"':
                state = "quoted"
            elif char in delimiter + "\r\n":
                state = "start"
            else:
                raise InputError("malformed_csv", "Unexpected text after a quoted field.")
        elif char == '"':
            if state != "start":
                raise InputError("malformed_csv", "Quotes must enclose the entire field.")
            state = "quoted"
        elif char in delimiter + "\r\n":
            state = "start"
        else:
            state = "unquoted"
    if state == "quoted":
        raise InputError("malformed_csv", "A quoted field is not closed.")


def _text_rows(content: bytes, delimiter: str) -> Iterable[list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise InputError("unsupported_encoding", "CSV and TSV must use UTF-8, with an optional BOM.", 415) from exc
    if not text.strip():
        raise InputError("empty_file", "The file is empty.")
    if "\x00" in text:
        raise InputError("unsupported_encoding", "CSV must not contain NUL characters.", 415)
    _check_quoting(text, delimiter)
    try:
        yield from csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    except csv.Error as exc:
        raise InputError("malformed_csv", "Delimited text quoting or record syntax is invalid.") from exc


def _bounds(rows: int, columns: int, settings: Settings) -> None:
    for value, maximum, code in ((rows, settings.max_rows, "too_many_rows"),
                                  (columns, settings.max_columns, "too_many_columns"),
                                  (rows * columns, settings.max_cells, "too_many_cells")):
        if value > maximum:
            raise InputError(code, "The dataset exceeds the configured size limit.", 413)


def _decoded_limit(size: int, settings: Settings) -> None:
    if size > settings.max_decoded_bytes:
        raise InputError("decoded_too_large", "The expanded dataset exceeds the configured byte limit.", 413)


def scalar_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, Decimal)):
        return str(value)
    raise InputError("unsupported_cell", "Only scalar text, numeric, boolean, and date/time values are supported.")


def _worksheet(book: Any, sheet_name: str | None) -> Any:
    names = [sheet.title for sheet in book.worksheets]
    if sheet_name is None:
        if len(names) != 1:
            raise InputError("sheet_required", "Select a worksheet using sheet_name.", details={"sheets": names})
        sheet_name = names[0]
    if sheet_name not in names:
        raise InputError("unknown_sheet", "The selected worksheet does not exist.", details={"sheets": names})
    return book[sheet_name]


def _xlsx_rows(sheet: Any, settings: Settings) -> Iterable[list[str]]:
    _bounds(max((sheet.max_row or 1) - 1, 0), sheet.max_column or 0, settings)
    # Ignore unreliable saved dimensions; inspect every actual row, including sparse rows.
    sheet.reset_dimensions()
    width = 0
    for index, cells in enumerate(sheet.iter_rows()):
        _bounds(index, len(cells), settings)
        if index == 0:
            if any(cell.value is not None and not isinstance(cell.value, str) for cell in cells):
                raise InputError("invalid_header", "Worksheet headers must be text.")
            width = len(cells)
        values: list[str] = []
        for cell in cells:
            if cell.data_type in ("f", "e"):
                raise InputError("unsupported_cell", "The selected worksheet must contain values, without formulas or errors.")
            value = cell.value
            if isinstance(value, datetime) and is_datetime(cell.number_format) == "date":
                value = value.date()
            values.append(scalar_text(value))
        if index and len(values) < width:
            values.extend([""] * (width - len(values)))
        yield values


def _supported_arrow_type(kind: Any) -> bool:
    if pa.types.is_dictionary(kind):
        return _supported_arrow_type(kind.value_type)
    return bool(pa.types.is_string(kind) or pa.types.is_large_string(kind) or pa.types.is_boolean(kind)
                or pa.types.is_integer(kind) or pa.types.is_floating(kind) or pa.types.is_decimal(kind)
                or pa.types.is_date(kind) or pa.types.is_timestamp(kind) or pa.types.is_time(kind)
                or pa.types.is_null(kind))


def _parquet_rows(content: bytes, settings: Settings) -> Iterable[list[str]]:
    with pq.ParquetFile(io.BytesIO(content), thrift_string_size_limit=settings.max_decoded_bytes,
                        thrift_container_size_limit=settings.max_cells) as file:
        _bounds(file.metadata.num_rows, len(file.schema_arrow), settings)
        if any(not _supported_arrow_type(field.type) for field in file.schema_arrow):
            raise InputError("unsupported_column_type", "Parquet requires flat scalar columns; nested and binary columns are unsupported.")
        expanded = sum(file.metadata.row_group(group).column(column).total_uncompressed_size
                       for group in range(file.metadata.num_row_groups)
                       for column in range(file.metadata.num_columns))
        _decoded_limit(expanded, settings)
        yield list(file.schema_arrow.names)
        for batch in file.iter_batches(batch_size=1024, use_threads=False):
            for index in range(batch.num_rows):
                yield [scalar_text(column[index].as_py()) for column in batch.columns]


def _collect(source: Iterable[list[str]], settings: Settings) -> tuple[list[str], tuple[tuple[str, ...], ...]]:
    iterator = iter(source)
    headers = next(iterator, [])
    if not headers or any(is_missing(header) for header in headers):
        raise InputError("blank_header", "Every column must have a nonblank header.")
    if len(set(headers)) != len(headers):
        raise InputError("duplicate_header", "Column headers must be unique.")
    _bounds(0, len(headers), settings)
    rows: list[tuple[str, ...]] = []
    size = sum(len(value.encode("utf-8")) for value in headers)
    for number, row in enumerate(iterator, 1):
        _bounds(number, len(headers), settings)
        if len(row) != len(headers):
            raise InputError("malformed_row", "A record has a different number of fields than the header.",
                             details={"record_number": number, "expected_fields": len(headers), "actual_fields": len(row)})
        if any("\x00" in value for value in row):
            raise InputError("unsupported_cell", "Data values must not contain NUL characters.")
        size += sum(len(value.encode("utf-8")) for value in row)
        _decoded_limit(size, settings)
        rows.append(tuple(row))
    if not rows:
        raise InputError("zero_rows", "The file contains a header but no data records.")
    return headers, tuple(rows)


def parse_dataset(content: bytes, filename: str | None, settings: Settings,
                  content_type: str | None = None, sheet_name: str | None = None) -> ParsedDataset:
    original_filename = validate_file_type(filename, content_type)
    file_format = cast(DataFormat, PurePath(original_filename).suffix.lower()[1:])
    if len(content) > settings.max_upload_bytes:
        raise InputError("upload_too_large", "The file exceeds the upload byte limit.", 413)
    if sheet_name is not None and file_format != "xlsx":
        raise InputError("unexpected_sheet", "sheet_name applies only to Excel files.")
    try:
        if file_format == "xlsx":
            with ZipFile(io.BytesIO(content)) as archive:
                _decoded_limit(sum(entry.file_size for entry in archive.infolist()), settings)
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            try:
                sheet = _worksheet(book, sheet_name)
                sheet_name = sheet.title
                headers, rows = _collect(_xlsx_rows(sheet, settings), settings)
            finally:
                book.close()
        elif file_format == "parquet":
            headers, rows = _collect(_parquet_rows(content, settings), settings)
        else:
            headers, rows = _collect(_text_rows(content, "\t" if file_format == "tsv" else ","), settings)
    except InputError:
        raise
    except Exception as exc:
        raise InputError("malformed_file", f"The {file_format} file could not be read.") from exc
    metadata = DatasetMetadata(
        original_filename=original_filename, content_sha256=hashlib.sha256(content).hexdigest(),
        size_bytes=len(content), row_count=len(rows), column_count=len(headers), headers=headers,
        format=file_format, encoding="utf-8" if file_format in ("csv", "tsv") else None, sheet_name=sheet_name,
    )
    return ParsedDataset(metadata=metadata, records=rows)


def replace_values(source: bytes, parsed: ParsedDataset, column: int, records: set[int],
                   replacement: str, settings: Settings) -> tuple[bytes, ParsedDataset]:
    expected = tuple(tuple(replacement if number in records and index == column else value
                           for index, value in enumerate(row)) for number, row in enumerate(parsed.records, 1))
    file_format = parsed.metadata.format
    buffer = io.BytesIO()
    try:
        if file_format == "xlsx":
            book = openpyxl.load_workbook(io.BytesIO(source))
            try:
                sheet = _worksheet(book, parsed.metadata.sheet_name)
                for number in records:
                    cell = sheet.cell(number + 1, column + 1)
                    cell.value = replacement
                    cell.data_type = "s"
                book.save(buffer)
            finally:
                book.close()
        elif file_format == "parquet":
            table = pq.read_table(io.BytesIO(source))
            field = table.schema.field(column)
            kind = field.type.value_type if pa.types.is_dictionary(field.type) else field.type
            if pa.types.is_null(kind):
                kind = pa.string()
            value = None if replacement == "" and not (pa.types.is_string(kind) or pa.types.is_large_string(kind)) else replacement
            scalar = pa.scalar(value).cast(kind)
            if scalar_text(scalar.as_py()) != replacement or (value is None and not field.nullable):
                raise ValueError("Replacement cannot be represented exactly")
            values = table.column(column).to_pylist()
            for number in records:
                values[number - 1] = scalar.as_py()
            target_type = kind if pa.types.is_null(field.type) else field.type
            updated = pa.array(values, type=target_type)
            table = table.set_column(column, field.with_type(target_type), updated)
            pq.write_table(table, buffer)
        else:
            text_buffer = io.StringIO(newline="")
            writer = csv.writer(text_buffer, delimiter="\t" if file_format == "tsv" else ",", lineterminator="\r\n")
            writer.writerow(parsed.metadata.headers)
            writer.writerows(expected)
            buffer.write(text_buffer.getvalue().encode("utf-8"))
    except Exception as exc:
        raise InputError("incompatible_replacement", "The replacement cannot be stored exactly in this file format and column type.") from exc
    content = buffer.getvalue()
    corrected = parse_dataset(content, parsed.metadata.original_filename, settings, sheet_name=parsed.metadata.sheet_name)
    if corrected.metadata.headers != parsed.metadata.headers or corrected.records != expected:
        raise InputError("incompatible_replacement", "The file format would change values outside the requested correction.")
    return content, corrected

